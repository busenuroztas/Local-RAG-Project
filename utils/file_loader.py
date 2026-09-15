import os
from pypdf import PdfReader

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

def load_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    
    try:
        if ext == ".pdf":
            reader = PdfReader(file_path)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
                    
        elif ext == ".docx":
            if DOCX_AVAILABLE:
                doc = Document(file_path)
                for paragraph in doc.paragraphs:
                    if paragraph.text:
                        text += paragraph.text + "\n"
            else:
                print("'python-docx' kütüphanesi yüklü değil.")

        elif ext == ".doc":
            if DOCX_AVAILABLE:
                try:
                    doc = Document(file_path)
                    for paragraph in doc.paragraphs:
                        if paragraph.text:
                            text += paragraph.text + "\n"
                except Exception:
                    pass
            if not text:
                with open(file_path, "rb") as f:
                    content = f.read()
                    text = "".join([chr(b) for b in content if 32 <= b <= 126 or b in (10, 13, 9)])
                    
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
                
    except Exception as e:
        print(f"Dosya okuma hatası ({file_path}): {e}")
        
    return text

def chunk_text(text, chunk_size=1000, overlap=150):
    
    if not text or not text.strip():
        return []

    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks = []
    current_chunk = ""

    for paragraph in paragraphs:
        if len(current_chunk) + len(paragraph) <= chunk_size:
            current_chunk += paragraph + "\n"
        else:
            if current_chunk:
                chunks.append(current_chunk.strip())
            current_chunk = paragraph + "\n"

    if current_chunk:
        chunks.append(current_chunk.strip())

    return chunks